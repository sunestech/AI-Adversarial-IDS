from pathlib import Path
import re
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_MD = PROJECT_ROOT / "reports" / "final_report.md"
REPORT_DOCX = PROJECT_ROOT / "reports" / "final_report.docx"

if not REPORT_MD.exists():
    raise FileNotFoundError(f"Missing report file: {REPORT_MD}")

text = REPORT_MD.read_text(encoding="utf-8")

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = None

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Adversarial Machine Learning Assessment of an IDS Classifier")
run.bold = True
run.font.size = None

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Security Research / Red-Team Style Report")

doc.add_paragraph("")

def clean_inline(s):
    s = s.strip()
    s = s.replace("**", "")
    s = s.replace("`", "")
    return s

def is_table_line(line):
    return line.strip().startswith("|") and line.strip().endswith("|")

def parse_table(lines):
    rows = []
    for line in lines:
        cells = [clean_inline(c) for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r"[-:\s]+", c) for c in cells):
            continue
        rows.append(cells)
    return rows

lines = text.splitlines()
i = 0

while i < len(lines):
    line = lines[i].rstrip()

    if not line.strip():
        i += 1
        continue

    if line.startswith("# "):
        # Skip duplicate markdown title because we already added a formatted title.
        i += 1
        continue

    if line.startswith("## "):
        doc.add_heading(clean_inline(line[3:]), level=1)
        i += 1
        continue

    if line.startswith("### "):
        doc.add_heading(clean_inline(line[4:]), level=2)
        i += 1
        continue

    if line.startswith("- "):
        doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
        i += 1
        continue

    if is_table_line(line):
        table_lines = []
        while i < len(lines) and is_table_line(lines[i]):
            table_lines.append(lines[i])
            i += 1

        rows = parse_table(table_lines)
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"

            for r_idx, row in enumerate(rows):
                for c_idx, cell_text in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_text
                    if r_idx == 0:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
            doc.add_paragraph("")
        continue

    p = doc.add_paragraph(clean_inline(line))
    i += 1

doc.add_page_break()

appendix = doc.add_heading("Appendix: Key Output Files", level=1)
for item in [
    "reports/final_report.md",
    "reports/final_report.docx",
    "reports/final_report.pdf",
    "results/metrics/final_defense_comparison_table.csv",
    "results/metrics/baseline_metrics.json",
    "results/metrics/feature_feasibility_table.csv",
    "results/metrics/phase3_whitebox_full_summary.csv",
    "results/metrics/phase4a_transfer_full_summary.csv",
    "results/metrics/phase4b_hopskipjump_full_summary.csv",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.save(REPORT_DOCX)

print("DOCX report saved:")
print(REPORT_DOCX)
